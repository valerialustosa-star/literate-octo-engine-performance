"""
PAINEL DE PREMIACAO - ASSIDUIDADE DA EQUIPE
============================================
Este programa foi feito para ser simples de usar:
1. Voce cola este codigo no Streamlit Community Cloud (gratuito).
2. Abre o site que ele gerar.
3. Faz upload da planilha "ABS AGOSTO.xlsx" (ou de outro mes).
4. O painel mostra automaticamente quem ganhou a premiacao do mes
   (quem teve o MENOR ABS = menor indice de ausencias/atrasos).

Nao precisa mudar nada no codigo todo mes: e so trocar a planilha
que voce sobe no site.
"""

import io
import datetime as dt

import streamlit as st
import pandas as pd
import plotly.express as px
import openpyxl


# ------------------------------------------------------------------
# CONFIGURACAO GERAL DA PAGINA (titulo na aba do navegador, layout largo)
# ------------------------------------------------------------------
st.set_page_config(
    page_title="Painel de Premiação - Assiduidade",
    page_icon="🏆",
    layout="wide",
)

# Um pouco de estilo visual para deixar parecido com um painel profissional
st.markdown(
    """
    <style>
        .metric-card {
            background-color: #ffffff;
            padding: 1.2rem;
            border-radius: 12px;
            box-shadow: 0 1px 4px rgba(0,0,0,0.12);
            text-align: center;
        }
        .winner-box {
            background: linear-gradient(135deg, #FFD700 0%, #FFA500 100%);
            padding: 1.5rem;
            border-radius: 16px;
            text-align: center;
            color: #202020;
        }
        .winner-box h1 {
            margin-bottom: 0;
        }
    </style>
    """,
    unsafe_allow_html=True,
)

st.title("🏆 Painel de Premiação - Assiduidade da Equipe")
st.caption(
    "Faça o upload da planilha de ABS (assiduidade) do mês e veja automaticamente "
    "o ranking e quem ganhou a premiação."
)


# ------------------------------------------------------------------
# FUNÇÕES DE APOIO
# ------------------------------------------------------------------

def timedelta_para_texto(valor):
    """Transforma um valor de tempo (horas trabalhadas) em um texto tipo '160h 30min'.
    Se não for um tempo válido, devolve um traço."""
    if isinstance(valor, dt.timedelta):
        total_minutos = int(valor.total_seconds() // 60)
        horas = total_minutos // 60
        minutos = total_minutos % 60
        return f"{horas}h {minutos:02d}min"
    return "-"


def encontrar_aba_e_colunas(caminho_arquivo):
    """Procura, dentro do arquivo Excel, a aba que tem uma coluna 'NOME' e
    uma coluna 'ABS'. Isso permite que o programa funcione mesmo que a aba
    não se chame exatamente 'GLOBAL', desde que tenha essas colunas.
    Devolve: nome_da_aba, dicionario_de_colunas (nome_da_coluna -> numero_da_coluna)
    """
    workbook = openpyxl.load_workbook(caminho_arquivo, data_only=True)

    # Prioriza abas com nomes comuns de resumo, mas verifica todas se precisar
    ordem_de_busca = list(workbook.sheetnames)
    if "GLOBAL" in ordem_de_busca:
        ordem_de_busca.remove("GLOBAL")
        ordem_de_busca.insert(0, "GLOBAL")

    for nome_aba in ordem_de_busca:
        aba = workbook[nome_aba]
        # Procura o cabeçalho nas primeiras 10 linhas da aba
        for linha in aba.iter_rows(min_row=1, max_row=min(10, aba.max_row)):
            colunas = {}
            for celula in linha:
                if isinstance(celula.value, str):
                    texto = celula.value.strip().upper()
                    if texto in ("NOME", "ABS", "PREVISTO", "REALIZADO",
                                 "TEMPO MÉDIO DE TRABALHO", "JUSTIFICATIVA",
                                 "AUSÊNCIAS"):
                        colunas[texto] = celula.column
            if "NOME" in colunas and "ABS" in colunas:
                return nome_aba, colunas, celula.row

    return None, None, None


@st.cache_data(show_spinner=False)
def carregar_dados(arquivo_bytes):
    """Lê o arquivo Excel enviado pelo usuário e devolve uma tabela (DataFrame)
    já organizada, pronta para ser exibida no painel."""

    caminho_em_memoria = io.BytesIO(arquivo_bytes)
    nome_aba, colunas, linha_cabecalho = encontrar_aba_e_colunas(caminho_em_memoria)

    if nome_aba is None:
        return None, None

    caminho_em_memoria.seek(0)
    workbook = openpyxl.load_workbook(caminho_em_memoria, data_only=True)
    aba = workbook[nome_aba]

    linhas = []
    for r in range(linha_cabecalho + 1, aba.max_row + 1):
        nome = aba.cell(r, colunas["NOME"]).value
        abs_valor = aba.cell(r, colunas["ABS"]).value

        # Pula linhas vazias, de total, ou sem um índice de ABS válido
        if not nome or not isinstance(nome, str):
            continue
        if abs_valor is None or not isinstance(abs_valor, (int, float)):
            continue

        previsto = aba.cell(r, colunas["PREVISTO"]).value if "PREVISTO" in colunas else None
        realizado = aba.cell(r, colunas["REALIZADO"]).value if "REALIZADO" in colunas else None
        tempo_medio = aba.cell(r, colunas["TEMPO MÉDIO DE TRABALHO"]).value if "TEMPO MÉDIO DE TRABALHO" in colunas else None
        justificativa = aba.cell(r, colunas["JUSTIFICATIVA"]).value if "JUSTIFICATIVA" in colunas else None
        ausencias = aba.cell(r, colunas["AUSÊNCIAS"]).value if "AUSÊNCIAS" in colunas else None

        linhas.append({
            "Nome": nome.strip(),
            "ABS (%)": round(abs_valor * 100, 2),
            "Horas Previstas": timedelta_para_texto(previsto),
            "Horas Realizadas": timedelta_para_texto(realizado),
            "Tempo Médio de Trabalho/Dia": timedelta_para_texto(tempo_medio),
            "Justificativa": justificativa if justificativa else "-",
            "Ausências": ausencias if ausencias else "-",
        })

    if not linhas:
        return None, nome_aba

    tabela = pd.DataFrame(linhas)
    tabela = tabela.sort_values(by="ABS (%)", ascending=True).reset_index(drop=True)
    tabela.insert(0, "Posição", range(1, len(tabela) + 1))
    return tabela, nome_aba


# ------------------------------------------------------------------
# UPLOAD DA PLANILHA
# ------------------------------------------------------------------
arquivo_enviado = st.file_uploader(
    "📂 Envie a planilha do mês (arquivo .xlsx)",
    type=["xlsx"],
)

if arquivo_enviado is None:
    st.info("Envie a planilha (ex: ABS AGOSTO.xlsx) para ver o painel.")
    st.stop()

tabela, aba_usada = carregar_dados(arquivo_enviado.getvalue())

if tabela is None:
    st.error(
        "Não consegui encontrar os dados de assiduidade nesta planilha. "
        "Verifique se ela tem uma aba com as colunas 'NOME' e 'ABS' "
        "(como a aba 'GLOBAL' do modelo original)."
    )
    st.stop()

st.success(f"Planilha carregada com sucesso! Dados lidos da aba **{aba_usada}**.")


# ------------------------------------------------------------------
# QUEM GANHOU A PREMIAÇÃO
# ------------------------------------------------------------------
menor_abs = tabela["ABS (%)"].min()
vencedores = tabela[tabela["ABS (%)"] == menor_abs]

st.markdown("### 🥇 Resultado da Premiação do Mês")

if len(vencedores) == 1:
    nome_vencedor = vencedores.iloc[0]["Nome"]
    st.markdown(
        f"""
        <div class="winner-box">
            <p style="margin-bottom:0.3rem;">🏆 VENCEDOR(A) DO MÊS 🏆</p>
            <h1>{nome_vencedor}</h1>
            <p>Índice de ABS: <strong>{menor_abs}%</strong></p>
        </div>
        """,
        unsafe_allow_html=True,
    )
else:
    nomes_empate = ", ".join(vencedores["Nome"].tolist())
    st.markdown(
        f"""
        <div class="winner-box">
            <p style="margin-bottom:0.3rem;">🏆 EMPATE NA PREMIAÇÃO 🏆</p>
            <h2>{nomes_empate}</h2>
            <p>Todos com o menor índice de ABS: <strong>{menor_abs}%</strong></p>
        </div>
        """,
        unsafe_allow_html=True,
    )

st.write("")

# ------------------------------------------------------------------
# CARDS COM NÚMEROS GERAIS DA EQUIPE
# ------------------------------------------------------------------
col1, col2, col3 = st.columns(3)
with col1:
    st.markdown(
        f'<div class="metric-card"><h3>{len(tabela)}</h3><p>Colaboradores</p></div>',
        unsafe_allow_html=True,
    )
with col2:
    media_abs = round(tabela["ABS (%)"].mean(), 2)
    st.markdown(
        f'<div class="metric-card"><h3>{media_abs}%</h3><p>Média de ABS da equipe</p></div>',
        unsafe_allow_html=True,
    )
with col3:
    st.markdown(
        f'<div class="metric-card"><h3>{menor_abs}%</h3><p>Menor ABS (vencedor)</p></div>',
        unsafe_allow_html=True,
    )

st.write("")

# ------------------------------------------------------------------
# GRÁFICO ESTILO POWER BI
# ------------------------------------------------------------------
st.markdown("### 📊 Ranking de Assiduidade (ABS por colaborador)")

tabela_grafico = tabela.copy()
tabela_grafico["Cor"] = tabela_grafico["ABS (%)"].apply(
    lambda x: "Vencedor(a)" if x == menor_abs else "Demais"
)

grafico = px.bar(
    tabela_grafico.sort_values("ABS (%)", ascending=True),
    x="ABS (%)",
    y="Nome",
    orientation="h",
    color="Cor",
    color_discrete_map={"Vencedor(a)": "#FFD700", "Demais": "#4C72B0"},
    text="ABS (%)",
)
grafico.update_layout(
    yaxis={"categoryorder": "total descending"},
    xaxis_title="ABS (%) — quanto menor, melhor",
    yaxis_title="",
    showlegend=False,
    height=max(400, 28 * len(tabela)),
)
grafico.update_traces(texttemplate="%{text}%", textposition="outside")

st.plotly_chart(grafico, use_container_width=True)


# ------------------------------------------------------------------
# TABELA COMPLETA COM FILTRO
# ------------------------------------------------------------------
st.markdown("### 📋 Tabela Completa")

busca = st.text_input("🔎 Filtrar por nome (opcional)")
tabela_exibicao = tabela.copy()
if busca:
    tabela_exibicao = tabela_exibicao[
        tabela_exibicao["Nome"].str.contains(busca, case=False, na=False)
    ]

st.dataframe(tabela_exibicao, use_container_width=True, hide_index=True)

# Botão para baixar o ranking em CSV, caso queira usar em outro lugar
csv = tabela.to_csv(index=False).encode("utf-8-sig")
st.download_button(
    "⬇️ Baixar ranking em CSV",
    data=csv,
    file_name="ranking_premiacao.csv",
    mime="text/csv",
)
